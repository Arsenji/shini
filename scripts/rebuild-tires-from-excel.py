#!/usr/bin/env python3
"""Rebuild passenger / lcv / truck tires from the three site price Excel files.

Preserves disk / tube / rimTape entries and remaps tire photos from the previous catalog.
"""

from __future__ import annotations

import json
import re
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
PRODUCTS_TS = ROOT / "src/data/shop/products.ts"
DOWNLOADS = Path.home() / "Downloads"

FILES = {
    "truck": DOWNLOADS / "Прайс для сайта (ГРУЗОВЫЕ).xlsx",
    "winter": DOWNLOADS / "Прайс для сайта (ЛЕГКОВЫЕ ЗИМА).xlsx",
    "summer": DOWNLOADS / "Прайс для сайта (ЛЕГКОВЫЕ ЛЕТО).xlsx",
}

SIZE_RE = re.compile(
    r"(?P<size>"
    # 205/55R16, 215/75R17.5, 215/75 R17.5
    r"\d{2,3}(?:[.,]\d+)?/\d{2,3}\s*[Zz]?[Rr]\s*\d{2}(?:[.,]\d+)?C?"
    # 8.25R20, 9.00 R20, 11R22.5, 12.00R20
    r"|\d{1,3}(?:[.,]\d+)?\s*[Zz]?[Rr]\s*\d{2}(?:[.,]\d+)?"
    # 6.50-16, 7.50-20, 11.2-20
    r"|\d{1,2}(?:[.,]\d+)?\s*[-xх×]\s*\d{2,3}(?:[.,]\d+)?"
    # 205/55-16
    r"|\d{2,3}/\d{2,3}\s*[-xх×]\s*\d{2}(?:[.,]\d+)?"
    r")",
    re.I,
)
STOCK_RE = re.compile(r"\(?\s*(\d+)\s*шт\.?\s*\)?", re.I)
# Stud / friction markers (order matters)
GRIP_STUD_RE = re.compile(
    r"(?<![А-Яа-яA-Za-z])(?:не\s*)?(?:шип\.?|шин\.?|ш\.|\(ш\))(?![А-Яа-яA-Za-z])",
    re.I,
)
GRIP_FRICTION_RE = re.compile(r"\bлип(?:учка|а|учки)?\.?\b", re.I)
NOT_STUD_RE = re.compile(r"\bне\s*шип\.?\b", re.I)
ASHINA_RE = re.compile(r"\b(?:а/?шина|автошина)\b", re.I)
HEADER_SKIP = re.compile(
    r"^(?:наименование|цена|лето|зима|всесезон|грузовые|спецтехника|сельхоз|"
    r"легковые|диски|камеры|лент)",
    re.I,
)


def normalize_size(raw: str) -> str:
    s = re.sub(r"\s+", "", raw.replace(",", ".").replace("\\", "/"))
    s = s.replace("х", "x").replace("×", "x").replace("Х", "x")
    s = re.sub(r"[Zz]?[Rr]", "R", s)
    # Repair common typo: 195/75R6C → 195/75R16C, 225/60R7 → 225/60R17
    s = re.sub(r"R([6-9])(C?)$", r"R1\1\2", s)
    if s.endswith("c"):
        s = s[:-1] + "C"
    s = re.sub(r"R(\d+)c$", r"R\1C", s)
    return s


def is_lcv_size(size: str) -> bool:
    return bool(re.search(r"C$", size, re.I))


def slugify(text: str) -> str:
    import hashlib

    t = text.lower().replace("\\", "/")
    t = re.sub(r"[х×]", "x", t)
    t = re.sub(r"\s+", "-", t.strip())
    ascii_part = re.sub(r"[^a-z0-9._-]+", "", t, flags=re.I)
    if ascii_part:
        return ascii_part[:48].strip("-")
    return hashlib.sha1(text.encode("utf-8")).hexdigest()[:10]


def extract_stock(name: str) -> tuple[str, int | None]:
    m = STOCK_RE.search(name)
    if not m:
        return name, None
    stock = int(m.group(1))
    cleaned = STOCK_RE.sub(" ", name)
    cleaned = re.sub(r"\s{2,}", " ", cleaned).strip(" ,;-")
    return cleaned, stock


def extract_grip(name: str) -> tuple[str, str | None]:
    """Return cleaned name and grip label: 'Шипы' | 'Липучка' | None."""
    if NOT_STUD_RE.search(name):
        cleaned = NOT_STUD_RE.sub(" ", name)
        cleaned = re.sub(r"\s{2,}", " ", cleaned).strip(" ,;-")
        return cleaned, None
    grip: str | None = None
    if GRIP_FRICTION_RE.search(name):
        grip = "Липучка"
        name = GRIP_FRICTION_RE.sub(" ", name)
    elif GRIP_STUD_RE.search(name):
        grip = "Шипы"
        name = GRIP_STUD_RE.sub(" ", name)
    cleaned = re.sub(r"\s{2,}", " ", name).strip(" ,;-")
    return cleaned, grip


def strip_ashina(name: str) -> str:
    return re.sub(r"\s{2,}", " ", ASHINA_RE.sub(" ", name)).strip(" ,;-")


def looks_like_size_header(name: str) -> bool:
    compact = re.sub(r"\s+", "", name)
    return bool(SIZE_RE.fullmatch(name)) or bool(
        re.fullmatch(
            r"\d{2,3}(?:[.,]\d+)?/\d{2,3}R\d{2}(?:[.,]\d+)?C?|\d{2,3}R\d{2}C?",
            compact,
            flags=re.I,
        )
    )


def parse_brand_model(rest: str) -> tuple[str, str]:
    tokens = rest.split()
    if not tokens:
        return "Unknown", rest
    # Multi-word brands
    two = " ".join(tokens[:2]).lower() if len(tokens) >= 2 else ""
    multi = {
        "double star",
        "road stone",
        "hi fly",
        "three-a",
        "three a",
    }
    if two in multi or (len(tokens) >= 2 and tokens[0].lower() in {"mp", "hk", "nf", "nr"}):
        brand = " ".join(tokens[:2])
        model = " ".join(tokens[2:]) or brand
        return brand, model
    brand = tokens[0]
    model = " ".join(tokens[1:]) if len(tokens) > 1 else tokens[0]
    return brand, model


def parse_passenger_line(name: str) -> tuple[str, str, str]:
    """size, brand, model from a passenger/lcv price line."""
    name = " ".join(name.replace("\\", "/").split())
    # Common typos in price lists: R6C → R16C, R7 → R17
    name = re.sub(r"(?<=[Rr])([6-9])(?=[Cc]|\b)", r"1\1", name)
    m = SIZE_RE.search(name)
    if not m:
        # brand-first without clear size — fail soft
        brand, model = parse_brand_model(name)
        return "", brand, model
    size = normalize_size(m.group("size"))
    before = name[: m.start()].strip(" -–—/")
    after = name[m.end() :].strip(" -–—/")
    if before and after:
        brand = before.split()[0] if len(before.split()) == 1 else before
        # Prefer brand before size when it's a known-looking brand chunk
        if len(before.split()) <= 3:
            brand = before
            model = after
        else:
            brand, model = parse_brand_model(f"{before} {after}")
    elif before:
        brand, model = parse_brand_model(before)
        if model == brand and after:
            model = after
        elif after:
            model = f"{model} {after}".strip()
    else:
        brand, model = parse_brand_model(after or name)
    model = model or size
    return size, brand.strip(), model.strip()


def parse_truck_line(name: str) -> tuple[str, str, str, str]:
    """Reuse truck parsing: size, brand, model, specs."""
    name = " ".join(name.replace("\\", "/").split())
    name = re.sub(
        rf"^({SIZE_RE.pattern})-(\d{{1,2}})\s+",
        r"\1 PR\2 ",
        name,
        flags=re.I,
    )

    def split_rest(rest: str, size: str) -> tuple[str, str, str]:
        tokens = rest.split()
        if not tokens:
            return size, size, ""
        brand = tokens[0]
        # Join "О 40" / "ИН 142" style codes broken by spaces
        if (
            len(tokens) >= 2
            and re.fullmatch(r"[A-Za-zА-Яа-яЁё]{1,3}", brand)
            and re.match(r"^\d", tokens[1])
        ):
            brand = f"{brand}-{tokens[1]}"
            tokens = [brand, *tokens[2:]]
        if len(tokens) >= 2:
            model = tokens[1]
            specs = " ".join(tokens[2:])
        else:
            model = tokens[0]
            specs = ""
        return brand, model, specs.strip()

    # size at start: 8.25R20 VM-201 ...
    m = re.match(rf"^{SIZE_RE.pattern}\s+(.+)$", name, re.I)
    if m:
        size = normalize_size(m.group("size"))
        rest = m.group(m.lastindex).strip()
        brand, model, specs = split_rest(rest, size)
        return size, brand, model, specs

    # size in the middle: У-2 8.25R20 Н/К
    m = re.match(rf"^(.+?)\s+{SIZE_RE.pattern}\s*(.*)$", name, re.I)
    if m:
        brand = m.group(1).strip()
        size = normalize_size(m.group("size"))
        tail = m.group(m.lastindex).strip()
        toks = tail.split()
        if toks:
            model = toks[0]
            specs = " ".join(toks[1:])
        else:
            model = brand
            specs = ""
        return size, brand, model, specs

    # dash sizes without R: 6.50-16 ...
    m = re.match(r"^(\d{1,2}(?:[.,]\d+)?-\d{2}(?:[.,]\d+)?)\s+(.+)$", name, re.I)
    if m:
        size = normalize_size(m.group(1))
        rest = m.group(2).strip()
        brand, model, specs = split_rest(rest, size)
        return size, brand, model, specs

    return name, name, name, ""


def extract_ply_rating(text: str) -> str | None:
    m = re.search(r"\b(\d{1,2})\s*сл\b", text, re.I)
    if m:
        return f"{m.group(1)} слоев"
    m = re.search(r"\b(?:PR\s*(\d{1,2})|(\d{1,2})\s*PR)\b", text, re.I)
    pr = m.group(1) if m and m.group(1) else (m.group(2) if m else None)
    if pr:
        return f"{pr} PR"
    return None


def extract_load_index(text: str) -> str | None:
    m = re.search(r"\b(\d{2,3}(?:/\d{2,3})?[A-ZА-Я])\b", text, re.I)
    if not m:
        return None
    return m.group(1).upper().replace("К", "K").replace("Т", "T").replace("Н", "H")


def load_workbook_rows(xlsx: Path) -> list[tuple[str, int]]:
    try:
        import openpyxl
    except ImportError:
        subprocess.check_call(
            [sys.executable, "-m", "pip", "install", "--trusted-host", "pypi.org",
             "--trusted-host", "files.pythonhosted.org", "openpyxl", "-q"]
        )
        import openpyxl

    wb = openpyxl.load_workbook(xlsx, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows: list[tuple[str, int]] = []
    for r in range(1, ws.max_row + 1):
        cell = ws.cell(r, 1).value
        price = ws.cell(r, 2).value
        if not isinstance(cell, str):
            continue
        name = cell.strip()
        if not name:
            continue
        if HEADER_SKIP.match(name) and not isinstance(price, (int, float)):
            continue
        if looks_like_size_header(name) and not isinstance(price, (int, float)):
            continue
        if not isinstance(price, (int, float)):
            continue
        rows.append((name, int(price)))
    wb.close()
    return rows


def make_id(parts: list[str], used: set[str]) -> str:
    base = "-".join(p for p in (slugify(x) for x in parts) if p)[:60].strip("-") or "tire"
    candidate = base
    n = 2
    while candidate in used:
        candidate = f"{base}-{n}"
        n += 1
    used.add(candidate)
    return candidate


def dump_existing_catalog() -> list[dict]:
    script = (
        "import { shopProducts } from './src/data/shop/products.ts'; "
        "import fs from 'fs'; "
        "fs.writeFileSync('.tmp-catalog.json', JSON.stringify(shopProducts))"
    )
    subprocess.check_call(
        ["node", "--import", "tsx", "-e", script],
        cwd=ROOT,
    )
    return json.loads((ROOT / ".tmp-catalog.json").read_text(encoding="utf-8"))


def build_image_index(existing: list[dict]) -> dict[str, str]:
    """Map brand|model keys → image path."""
    index: dict[str, str] = {}
    for p in existing:
        if p.get("category") not in {"passenger", "lcv", "truck"}:
            continue
        img = p.get("image")
        if not img:
            continue
        brand = str(p.get("brand") or "").strip().lower()
        model = str(p.get("model") or "").strip().lower()
        keys = [
            f"{brand}|{model}",
            f"{brand}|{slugify(model)}",
            f"{brand}|{model.split()[0]}" if model else "",
        ]
        for k in keys:
            if k and k not in index:
                index[k] = img
    return index


def find_image(index: dict[str, str], brand: str, model: str) -> str | None:
    b = brand.strip().lower()
    m = model.strip().lower()
    for k in (f"{b}|{m}", f"{b}|{slugify(m)}", f"{b}|{m.split()[0]}" if m else ""):
        if k and k in index:
            return index[k]
    # looser: any key starting with brand|
    prefix = f"{b}|"
    for k, img in index.items():
        if k.startswith(prefix) and (slugify(m) in k or (m and m.split()[0] in k)):
            return img
    return None


def build_tire_products(image_index: dict[str, str]) -> list[dict]:
    used: set[str] = set()
    products: list[dict] = []

    # --- Summer + Winter passenger/lcv ---
    for season, path in (("summer", FILES["summer"]), ("winter", FILES["winter"])):
        for raw, price in load_workbook_rows(path):
            name, stock = extract_stock(raw)
            name, grip = extract_grip(name)
            name = strip_ashina(name)
            size, brand, model = parse_passenger_line(name)
            if not size:
                # try again on original without grip strip issues
                size, brand, model = parse_passenger_line(strip_ashina(extract_stock(raw)[0]))
            if not size:
                print(f"WARN skip (no size): {raw!r}")
                continue
            category = "lcv" if is_lcv_size(size) else "passenger"
            # Append grip words into model for searchability if not already
            model_clean = model
            if grip == "Шипы" and not re.search(r"шип", model_clean, re.I):
                model_clean = f"{model_clean} шипы".strip()
            if grip == "Липучка" and not re.search(r"лип", model_clean, re.I):
                model_clean = f"{model_clean} липучка".strip()

            pid = make_id([brand, model_clean, size, season], used)
            item: dict = {
                "id": pid,
                "brand": brand,
                "model": model_clean,
                "category": category,
                "season": season,
                "imageKey": category,
                "sizeGroup": size,
                "sizes": [size],
                "price": price,
                "offers": [{"size": size, "price": price}],
            }
            if grip:
                item["badge"] = grip
            if stock:
                item["stockRemaining"] = stock
            img = find_image(image_index, brand, model_clean) or find_image(image_index, brand, model)
            if img:
                item["image"] = img
            load = extract_load_index(f"{model} {raw}")
            if load:
                item["loadIndex"] = load
            products.append(item)

    # --- Truck / special / agro (all in freight file) ---
    for raw, price in load_workbook_rows(FILES["truck"]):
        name, stock = extract_stock(raw)
        name, grip = extract_grip(name)
        name = strip_ashina(name)
        size, brand, model, specs = parse_truck_line(name)
        specs = strip_ashina(specs)
        if grip == "Шипы":
            specs = f"{specs} шипы".strip()
        elif grip == "Липучка":
            specs = f"{specs} липучка".strip()
        pid = make_id([brand, model, size, specs], used)
        item = {
            "id": pid,
            "brand": brand,
            "model": model,
            "category": "truck",
            "imageKey": "truck",
            "sizeGroup": size,
            "sizes": [size],
            "price": price,
            "offers": [{"size": size, "price": price}],
        }
        if specs:
            item["truckSpecs"] = specs
        if grip:
            item["badge"] = grip
        if stock:
            item["stockRemaining"] = stock
        full = " ".join([brand, model, specs, raw])
        ply = extract_ply_rating(full)
        load = extract_load_index(full)
        if ply:
            item["plyRating"] = ply
        if load:
            item["loadIndex"] = load
        img = find_image(image_index, brand, model)
        if img:
            item["image"] = img
        products.append(item)

    return products


def ts_str(value: object) -> str:
    if value is None:
        return "null"
    return "'" + str(value).replace("\\", "\\\\").replace("'", "\\'") + "'"


def render_product(p: dict) -> str:
    lines = [
        "  {",
        f"    id: {ts_str(p['id'])},",
        f"    brand: {ts_str(p['brand'])},",
        f"    model: {ts_str(p['model'])},",
        f"    category: {ts_str(p['category'])},",
    ]
    if p.get("season"):
        lines.append(f"    season: {ts_str(p['season'])},")
    lines.append(f"    imageKey: {ts_str(p['imageKey'])},")
    if p.get("sizeGroup"):
        lines.append(f"    sizeGroup: {ts_str(p['sizeGroup'])},")
    if p.get("sizes"):
        lines.append(f"    sizes: [{', '.join(ts_str(s) for s in p['sizes'])}],")
    if p.get("badge"):
        lines.append(f"    badge: {ts_str(p['badge'])},")
    if p.get("stockRemaining"):
        lines.append(f"    stockRemaining: {int(p['stockRemaining'])},")
    if p.get("image"):
        lines.append(f"    image: {ts_str(p['image'])},")
    if p.get("color"):
        lines.append(f"    color: {ts_str(p['color'])},")
    if p.get("truckSpecs"):
        lines.append(f"    truckSpecs: {ts_str(p['truckSpecs'])},")
    if p.get("plyRating"):
        lines.append(f"    plyRating: {ts_str(p['plyRating'])},")
    if p.get("loadIndex"):
        lines.append(f"    loadIndex: {ts_str(p['loadIndex'])},")
    if isinstance(p.get("price"), (int, float)):
        lines.append(f"    price: {int(p['price'])},")
    if p.get("offers"):
        lines.append("    offers: [")
        for o in p["offers"]:
            lines.append(f"      {{ size: {ts_str(o['size'])}, price: {int(o['price'])} }},")
        lines.append("    ],")
    lines.append("  }")
    return "\n".join(lines)


def write_products_ts(kept: list[dict], tires: list[dict]) -> None:
    header = """import type { ShopProduct } from './types'

/**
 * Каталог магазина.
 * Шины (passenger/lcv/truck) — автогенерация: scripts/rebuild-tires-from-excel.py
 * Диски / камеры / ободные ленты сохраняются из предыдущего каталога.
 */
export const shopProducts: ShopProduct[] = [
"""
    blocks = [render_product(p) for p in tires + kept]
    PRODUCTS_TS.write_text(header + ",\n".join(blocks) + "\n]\n", encoding="utf-8")


def main() -> None:
    for key, path in FILES.items():
        if not path.is_file():
            print(f"Missing Excel ({key}): {path}", file=sys.stderr)
            sys.exit(1)

    print("Loading existing catalog…")
    existing = dump_existing_catalog()
    kept = [p for p in existing if p.get("category") in {"disk", "tube", "rimTape"}]
    image_index = build_image_index(existing)
    print(f"  keep {len(kept)} disk/tube/rimTape; image keys {len(image_index)}")

    print("Parsing Excel…")
    tires = build_tire_products(image_index)
    by_cat: dict[str, int] = {}
    by_season: dict[str, int] = {}
    stock_n = grip_n = 0
    for p in tires:
        by_cat[p["category"]] = by_cat.get(p["category"], 0) + 1
        by_season[p.get("season") or "none"] = by_season.get(p.get("season") or "none", 0) + 1
        if p.get("stockRemaining"):
            stock_n += 1
        if p.get("badge") in {"Шипы", "Липучка"}:
            grip_n += 1
    print(f"  tires={len(tires)} by_cat={by_cat} by_season={by_season} stock={stock_n} grip={grip_n}")

    write_products_ts(kept, tires)
    print(f"Wrote {len(tires) + len(kept)} products → {PRODUCTS_TS}")


if __name__ == "__main__":
    main()
